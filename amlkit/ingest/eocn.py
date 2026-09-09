"""UAE Local Terrorist List -- direct ingestion from the EOCN.

WHY THIS ADAPTER EXISTS
-----------------------
The same list is available, already parsed, from OpenSanctions
(`opensanctions.py::uae_local_terrorists`). That adapter is licensed
CC-BY-NC 4.0: free for non-commercial use, with no exemption for commercial
users. This one goes to the primary source instead, so the mandatory UAE list
can be screened against by a product that is actually sold. That is the whole
point of the adapter boundary described in `base.py`.

WHAT THE SOURCE ACTUALLY IS
---------------------------
The Executive Office for Control & Non-proliferation publishes the list from
`/en-us/un-page` as two downloads of the same data -- a PDF and an Excel
workbook -- both served through the site's own file API by opaque
`FileID` GUIDs. The GUID changes when the Office uploads a revision, so it
cannot be hardcoded: `fetch()` re-reads the page every run and resolves the
current file, which is also how a new revision is picked up at all.

The workbook is chosen over the PDF deliberately. The PDF is a rendering of
the same table, and extracting it means reconstructing columns from text
positions -- a parser that fails silently and subtly when the layout shifts.
The workbook carries the table as a table.

WORKBOOK SHAPE
--------------
Five sheets, three of them listings and two of them delistings:

    الأفراد               individuals, currently listed
    التنظيمات             groups/organisations, currently listed
    الكيانات              legal entities (companies), currently listed
    رفع الإدراج - أفراد    individuals REMOVED from the list
    رفع الإدراج - كيانات   entities REMOVED from the list

The delisting sheets are skipped. They are the Office's record of who *used
to* be designated, and indexing them would generate alerts against people the
Cabinet has explicitly unlisted -- the same failure the loader's
replace-on-refresh semantics exist to prevent, reintroduced from inside the
source file. They are identified structurally (a `قرار رفع الإدراج`
-- "delisting resolution" -- column) rather than by sheet name, so a renamed
tab does not silently turn delisted people back into screening targets.

Every entry is emitted under programme `AE-UNSC1373`, the UNSCR 1373 basis for
the local list, which is what `screening/pf.py` matches on to classify a hit
as terrorism rather than proliferation financing.
"""

from __future__ import annotations

import io
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from typing import Iterator

from .base import AdapterError, SourceEntity

PAGE_URL = "https://www.uaeiec.gov.ae/en-us/un-page"
DOWNLOAD_URL = "https://www.uaeiec.gov.ae/API/Upload/DownloadFile?FileID={file_id}"
USER_AGENT = "amlkit/0.1 (UAE AML screening; compliance tooling)"

# UNSCR 1373 is the resolution under which the UAE maintains a national list;
# screening/pf.py keys terrorism-vs-proliferation classification off this.
PROGRAM = "AE-UNSC1373"

# An escape hatch, not a configuration knob. If the Office restructures the
# page faster than this adapter can be updated, a compliance officer can point
# it straight at the file rather than lose the mandatory list entirely.
ENV_URL_OVERRIDE = "AMLKIT_EOCN_LIST_URL"

_FILE_ID = r"DownloadFile\?FileID=(?P<fid>[0-9a-fA-F]{8}-(?:[0-9a-fA-F]{4}-){3}[0-9a-fA-F]{12})"

# One pass over the page, in document order, recognising four things: the
# heading of the list we want, the heading of the list that follows it, the
# little file-format icons the page puts beside each download, and the
# download links themselves. Scanning for all of them together is what lets
# `_candidate_file_ids` tell "the Excel file under Local Terrorist List" from
# "some other Excel file elsewhere on a page that publishes a dozen of them".
_MARKER = re.compile(
    r"(?P<ltl>Local\s+Terrorist\s+List)"
    r"|(?P<other>Consolidated\s+List)"
    r"|(?P<excel>icon-xlsx?[a-z-]*)"
    r"|(?P<pdf>icon-pdf[a-z-]*)"
    r"|" + _FILE_ID,
    re.IGNORECASE,
)

# How far past a "Local Terrorist List" heading a download link can appear and
# still plausibly belong to it. The page's own block is well under this.
_WINDOW_CHARS = 4000

# Downloads actually fetched while hunting for the workbook. The page offers
# the list in two formats, so anything beyond a handful means the structure
# has changed and guessing further is worse than failing loudly.
_MAX_DOWNLOAD_ATTEMPTS = 4

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"   # legacy .xls
_ZIP_MAGIC = b"PK\x03\x04"                          # .xlsx (a zip container)


# --------------------------------------------------------------------------
# Arabic text folding
# --------------------------------------------------------------------------

_DIACRITICS = re.compile("[\u064B-\u0652\u0670\u0640]")
_ALEF = re.compile(r"[أإآٱ]")


def _fold(value: object) -> str:
    """Normalise Arabic text for comparison: diacritics, alef and ya variants.

    The workbook is hand-maintained, so the same country or column heading
    appears with and without hamza, with ta marbuta or ha, occasionally with
    a stray tatweel. Folding those away is what lets header and country
    lookups be written once instead of once per spelling.
    """
    text = unicodedata.normalize("NFKC", str(value or ""))
    text = _DIACRITICS.sub("", text)
    text = _ALEF.sub("ا", text)
    text = text.replace("ى", "ي").replace("ة", "ه")
    text = text.replace("ؤ", "و").replace("ئ", "ي")
    return re.sub(r"\s+", " ", text).strip()


def _key(value: object) -> str:
    """Fold, then strip all whitespace -- for matching column headings.

    Headings in the workbook contain hard line breaks mid-phrase
    ("اسم العائلة\n (بالحروف اللاتينية)"), so whitespace cannot be part of
    the comparison.
    """
    return re.sub(r"\s+", "", _fold(value))


# Blank in this workbook is any of: empty, "-", "--", "N/A". Written out
# rather than inferred, because a literal "-" reaching the matcher as a name
# would be a token that matches everything.
_BLANKS = {"", "-", "--", "---", "n/a", "na", "غير معروف", "غير محدد"}


def _clean(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return None if text.lower() in _BLANKS else text


# --------------------------------------------------------------------------
# Country mapping
# --------------------------------------------------------------------------
#
# The nationality column is Arabic free text, and it is messy in ways that
# matter: adjectival forms ("لبناني" for Lebanon), spelling variants
# ("السعوديه"/"السعودية"), dual nationality on one line ("السعودية / الكويت"),
# and outright narrative ("الحالية: السويد السابقة: ليبيريا" -- current
# Sweden, former Liberia).
#
# Rather than tokenise that, every known country name is substring-matched
# against the folded cell, which handles all four shapes with one rule.
#
# Each match contributes BOTH the ISO code and the English name, and the raw
# Arabic string is kept too. That is deliberate: scorer.py applies a
# `country_mismatch` PENALTY when the customer's stated nationality is absent
# from the entity's country list, so a list holding only "قطر" would suppress
# a real hit on a customer recorded as "Qatar" or "QA". Under-populating this
# field costs true positives, which is the expensive direction to be wrong in.

_COUNTRIES: tuple[tuple[str, str, str], ...] = (
    # (Arabic form, ISO 3166-1 alpha-2, English name)
    ("قطر", "QA", "Qatar"),
    ("الاردن", "JO", "Jordan"),
    ("اليمن", "YE", "Yemen"),
    ("الكويت", "KW", "Kuwait"),
    ("ليبيا", "LY", "Libya"),
    ("ليبيريا", "LR", "Liberia"),
    ("مصر", "EG", "Egypt"),
    ("الامارات", "AE", "United Arab Emirates"),
    ("السعوديه", "SA", "Saudi Arabia"),
    ("البحرين", "BH", "Bahrain"),
    ("امريكا", "US", "United States"),
    ("الولايات المتحده", "US", "United States"),
    ("الصومال", "SO", "Somalia"),
    ("النمسا", "AT", "Austria"),
    ("ايران", "IR", "Iran"),
    ("افغانستان", "AF", "Afghanistan"),
    ("باكستان", "PK", "Pakistan"),
    ("سوريا", "SY", "Syria"),
    ("سوري", "SY", "Syria"),
    ("لبنان", "LB", "Lebanon"),
    ("نيجيريا", "NG", "Nigeria"),
    ("نيجيري", "NG", "Nigeria"),
    ("بريطانيا", "GB", "United Kingdom"),
    ("المملكه المتحده", "GB", "United Kingdom"),
    ("سانت كيتس", "KN", "Saint Kitts and Nevis"),
    ("روسيا", "RU", "Russia"),
    ("تركيا", "TR", "Turkey"),
    ("السويد", "SE", "Sweden"),
    ("العراق", "IQ", "Iraq"),
    ("تونس", "TN", "Tunisia"),
    ("الجزاير", "DZ", "Algeria"),
    ("المغرب", "MA", "Morocco"),
    ("السودان", "SD", "Sudan"),
    ("فلسطين", "PS", "Palestine"),
    ("عمان", "OM", "Oman"),
    ("تشاد", "TD", "Chad"),
    ("مالي", "ML", "Mali"),
    ("كينيا", "KE", "Kenya"),
    ("الهند", "IN", "India"),
    ("بنغلاديش", "BD", "Bangladesh"),
    ("اندونيسيا", "ID", "Indonesia"),
    ("ماليزيا", "MY", "Malaysia"),
    ("فرنسا", "FR", "France"),
    ("المانيا", "DE", "Germany"),
    ("بلجيكا", "BE", "Belgium"),
    ("هولندا", "NL", "Netherlands"),
)


def _countries(cell: object) -> list[str]:
    raw = _clean(cell)
    if not raw:
        return []
    folded = _fold(raw)
    out: list[str] = []
    for arabic, iso, english in _COUNTRIES:
        if arabic in folded:
            for form in (iso, english, arabic):
                if form not in out:
                    out.append(form)
    # Keep the original string even when nothing matched, so an unmapped
    # nationality still contributes to matching instead of vanishing -- and so
    # the gap is visible in the stored entity rather than only in this table.
    if raw not in out:
        out.append(raw)
    return out


# --------------------------------------------------------------------------
# Workbook reading
# --------------------------------------------------------------------------


@dataclass(slots=True)
class _Cell:
    """One cell, reduced to what the parser needs: a value and whether it is a date."""

    value: object
    is_date: bool = False


def _read_workbook(payload: bytes) -> list[tuple[str, list[list[_Cell]]]]:
    """Return [(sheet_name, rows)] for either workbook format.

    Both formats are supported because the source picking a different Excel
    flavour on some future upload must not take the mandatory UAE list
    offline. openpyxl and xlrd are both already dependencies.
    """
    if payload.startswith(_OLE2_MAGIC):
        return _read_xls(payload)
    if payload.startswith(_ZIP_MAGIC):
        return _read_xlsx(payload)
    raise AdapterError(
        "eocn: downloaded file is not an Excel workbook "
        f"(leading bytes {payload[:8]!r})"
    )


def _read_xls(payload: bytes) -> list[tuple[str, list[list[_Cell]]]]:
    import xlrd

    try:
        book = xlrd.open_workbook(file_contents=payload)
    except Exception as exc:  # xlrd raises a family of its own errors
        raise AdapterError(f"eocn: .xls parse failed - {exc}") from exc

    sheets = []
    for sheet in book.sheets():
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                is_date = sheet.cell_type(r, c) == xlrd.XL_CELL_DATE
                value = sheet.cell_value(r, c)
                if is_date:
                    value = _xldate(value, book.datemode)
                row.append(_Cell(value, is_date))
            rows.append(row)
        sheets.append((sheet.name, rows))
    return sheets


def _xldate(serial: float, datemode: int) -> str | None:
    """Excel serial -> ISO date. Returns None rather than raising on garbage."""
    import xlrd

    try:
        y, m, d, *_ = xlrd.xldate_as_tuple(serial, datemode)
        return date(y, m, d).isoformat()
    except Exception:
        return None


def _read_xlsx(payload: bytes) -> list[tuple[str, list[list[_Cell]]]]:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise AdapterError(f"eocn: .xlsx parse failed - {exc}") from exc

    sheets = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = []
            for value in row:
                if isinstance(value, datetime):
                    cells.append(_Cell(value.date().isoformat(), True))
                elif isinstance(value, date):
                    cells.append(_Cell(value.isoformat(), True))
                else:
                    cells.append(_Cell(value, False))
            rows.append(cells)
        sheets.append((sheet.title, rows))
    book.close()
    return sheets


# --------------------------------------------------------------------------
# Sheet interpretation
# --------------------------------------------------------------------------

# Column headings, folded. Aliases are listed where the workbook has used more
# than one wording across sheets.
H_INDEX = ("#",)
H_CLASS = ("التصنيف",)
H_NATIONALITY = ("الجنسيه",)
H_FAMILY_AR = ("اسمالعايله(بالحروفالعربيه)",)
H_FAMILY_LA = ("اسمالعايله(بالحروفاللاتينيه)",)
H_FULL_AR = ("الاسمالكامل(بالحروفالعربيه)", "الاسمالكامل(باللغهالعربيه)")
H_FULL_LA = ("الاسمالكامل(بالحروفاللاتينيه)", "الاسمالكامل(باللغهاللاتينيه)")
H_KNOWN_AS = ("الاسم",)
H_BIRTH_DATE = ("تاريخالميلاد",)
H_BIRTH_PLACE = ("مكانالميلاد",)
H_STREET = ("الشارع",)
H_CITY = ("المدينه",)
H_ADDR_COUNTRY = ("الدوله",)
H_DOC_TYPE = ("النوع",)
H_DOC_NUMBER = ("رقمالوثيقه",)
H_DOC_ISSUER = ("جههالاصدار",)
H_LICENCE_NO = ("رقمالرخصه",)
H_HQ = ("المقر",)
H_NOTES = ("ملاحظات",)
H_OTHER_INFO = ("معلوماتاخري",)
H_DELISTING = ("قراررفعالادراج",)

# التصنيف values. "شخص إرهابي" (terrorist person) is the only one that maps to
# a natural person; groups and legal entities both become Organization, which
# is what the matcher and the FtM vocabulary expect.
_PERSON_CLASS = "شخصارهابي"

_PASSPORT_HINT = "جوازسفر"
_PASSPORT_HINT_ALT = "جوازالسفر"

# Listing year, from the Arabic decision text: "... لسنة 2017".
_LISTED_YEAR = re.compile(r"لسنه\s*(\d{4})")


def _find_header(rows: list[list[_Cell]]) -> int | None:
    """Index of the header row -- the first whose first cell is '#'.

    The sheets open with one or two merged banner rows before the real
    headings, and how many varies by sheet.
    """
    for i, row in enumerate(rows[:10]):
        if row and _key(row[0].value) in H_INDEX:
            return i
    return None


def _columns(header: list[_Cell]) -> dict[str, int]:
    return {_key(cell.value): i for i, cell in enumerate(header) if _key(cell.value)}


def _col(columns: dict[str, int], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in columns:
            return columns[name]
    return None


def _at(row: list[_Cell], index: int | None) -> _Cell | None:
    if index is None or index >= len(row):
        return None
    return row[index]


def _text(row: list[_Cell], index: int | None) -> str | None:
    cell = _at(row, index)
    return _clean(cell.value) if cell else None


def _birth_date(cell: _Cell | None) -> str | None:
    """Birth dates arrive as three different things in the same column.

    A real date is stored as an Excel date (already ISO by the time it gets
    here). A year-only entry is stored as a bare number -- 1975, not a
    serial -- and must not be run through the serial conversion, which would
    silently turn it into 1905. Anything else is text.
    """
    if cell is None:
        return None
    if cell.is_date:
        return _clean(cell.value)
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if 1000 <= number <= 2200:
            return str(int(number))
        return None
    text = _clean(value)
    if not text:
        return None
    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    match = re.search(r"\b(1[89]\d{2}|20\d{2})\b", text)
    return match.group(1) if match else None


# Document and licence numbers are not always bare numbers. The workbook
# carries entries like "رقم الهوية اللبنانية :000010757716" and
# "رقم التسجيل (81396) رقم الرخصة (5083663376)" -- an Arabic label wrapped
# around one or two real numbers.
#
# Storing that whole string as the identifier would make it unmatchable:
# entity_identifiers is looked up by exact value, and no customer record will
# ever contain the label. So the numbers are pulled out individually and the
# original string is kept in `raw` for the evidence pack.
_ID_VALUE = re.compile(r"[A-Za-z0-9][A-Za-z0-9/\-]{2,}")


def _id_values(text: str | None) -> list[str]:
    if not text:
        return []
    out: list[str] = []
    for value in _ID_VALUE.findall(text):
        value = value.strip("/-")
        if value and value not in out:
            out.append(value)
    return out


def _identifiers(row: list[_Cell], columns: dict[str, int]) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    doc_type = _key(_text(row, _col(columns, H_DOC_TYPE)) or "")
    kind = "passport" if _PASSPORT_HINT in doc_type or _PASSPORT_HINT_ALT in doc_type else "id"
    for value in _id_values(_text(row, _col(columns, H_DOC_NUMBER))):
        out.append((kind, value))
    for value in _id_values(_text(row, _col(columns, H_LICENCE_NO))):
        out.append(("registration", value))
    return out


def _names(row: list[_Cell], columns: dict[str, int]) -> tuple[str | None, list[str]]:
    """Return (caption, all other names).

    The Latin full name is preferred as the caption because that is what an
    operator sees in an alert and what the printable evidence pack carries.
    The Arabic full name is never dropped -- Arabic-aware matching is this
    project's differentiator, and it works off the stored name rows.
    """
    latin = _text(row, _col(columns, H_FULL_LA))
    arabic = _text(row, _col(columns, H_FULL_AR))
    caption = latin or arabic
    names: list[str] = []
    for value in (
        arabic,
        latin,
        _text(row, _col(columns, H_KNOWN_AS)),
        _text(row, _col(columns, H_FAMILY_AR)),
        _text(row, _col(columns, H_FAMILY_LA)),
    ):
        if value and value != caption and value not in names:
            names.append(value)
    return caption, names


def _sheet_kind(columns: dict[str, int]) -> str:
    if _col(columns, H_NATIONALITY) is not None:
        return "IND"
    if _col(columns, H_LICENCE_NO) is not None:
        return "ENT"
    return "GRP"


class EOCNLocalTerroristListAdapter:
    """UAE Local Terrorist List, straight from the EOCN's own publication."""

    def __init__(self) -> None:
        # Deliberately the same dataset key the OpenSanctions-sourced adapter
        # used. upsert_dataset updates publisher and licence in place on
        # conflict, so an existing deployment swaps source without leaving a
        # second, permanently-stale UAE list behind in the staleness report.
        self.key = "ae_local_terrorists"
        self.title = "UAE Local Terrorist List"
        self.publisher = "Executive Office for Control & Non-proliferation (EOCN)"
        self.source_url = PAGE_URL
        # A UAE government publication that regulated entities are legally
        # required to screen against. Free to use, including commercially --
        # which is the entire reason this adapter exists.
        self.licence = "UAE Government publication - free to use"
        self.is_mandatory = True

    # -- fetch ------------------------------------------------------------

    def fetch(self) -> bytes:
        override = os.environ.get(ENV_URL_OVERRIDE, "").strip()
        if override:
            payload = self._download(override)
            if not _looks_like_workbook(payload):
                raise AdapterError(
                    f"{self.key}: {ENV_URL_OVERRIDE} does not point at an Excel workbook"
                )
            return payload

        page = self._download(PAGE_URL).decode("utf-8", errors="replace")
        file_ids = _candidate_file_ids(page)
        if not file_ids:
            raise AdapterError(
                f"{self.key}: no Local Terrorist List download found on {PAGE_URL} "
                f"- page structure may have changed (set {ENV_URL_OVERRIDE} to recover)"
            )

        seen_formats: list[str] = []
        for file_id in file_ids[:_MAX_DOWNLOAD_ATTEMPTS]:
            payload = self._download(DOWNLOAD_URL.format(file_id=file_id))
            if _looks_like_workbook(payload):
                return payload
            seen_formats.append(f"{file_id[:8]}={_describe(payload)}")

        raise AdapterError(
            f"{self.key}: Local Terrorist List is published but no Excel workbook "
            f"was among the downloads ({', '.join(seen_formats)}). The PDF is not "
            f"parsed by design - set {ENV_URL_OVERRIDE} to a workbook URL to recover."
        )

    def _download(self, url: str) -> bytes:
        from .base import fetch_with_retry
        return fetch_with_retry(self.key, url, user_agent=USER_AGENT)

    # -- parse ------------------------------------------------------------

    def parse(self, payload: bytes) -> Iterator[SourceEntity]:
        seen = 0
        skipped_sheets: list[str] = []

        for sheet_name, rows in _read_workbook(payload):
            header_index = _find_header(rows)
            if header_index is None:
                skipped_sheets.append(sheet_name)
                continue
            columns = _columns(rows[header_index])

            # A delisting sheet records people the Cabinet has REMOVED. Indexing
            # them would raise alerts against the unlisted. Detected by the
            # delisting-resolution column so a renamed tab cannot smuggle them
            # back into the screening index.
            if _col(columns, H_DELISTING) is not None:
                continue

            if _col(columns, H_FULL_LA) is None and _col(columns, H_FULL_AR) is None:
                skipped_sheets.append(sheet_name)
                continue

            kind = _sheet_kind(columns)
            for row in rows[header_index + 1:]:
                entity = _row_to_entity(row, columns, kind, sheet_name)
                if entity is not None:
                    seen += 1
                    yield entity

        if seen == 0:
            # Loud by design, same as every other adapter: a mandatory list
            # that silently parses to nothing looks identical to a healthy one
            # in the staleness report, which is the exact failure the 24-hour
            # rule exists to prevent.
            raise AdapterError(
                f"{self.key}: parsed 0 entities from the workbook "
                f"(sheets skipped: {skipped_sheets or 'none'}) - format may have changed"
            )


def _row_to_entity(
    row: list[_Cell], columns: dict[str, int], kind: str, sheet_name: str
) -> SourceEntity | None:
    index = _text(row, _col(columns, H_INDEX))
    caption, names = _names(row, columns)
    if not caption:
        return None
    # The workbook carries trailing formatting-only rows with a serial number
    # and nothing else; a row without a name is not an entry.
    if index is None:
        index = re.sub(r"\W+", "-", caption)[:40]

    classification = _text(row, _col(columns, H_CLASS))
    is_person = _key(classification or "") == _PERSON_CLASS or (
        classification is None and kind == "IND"
    )

    countries = _countries(_text(row, _col(columns, H_NATIONALITY)))
    for extra in (
        _text(row, _col(columns, H_ADDR_COUNTRY)),
        _text(row, _col(columns, H_HQ)),
    ):
        for value in _countries(extra):
            if value not in countries:
                countries.append(value)

    other_info = _text(row, _col(columns, H_OTHER_INFO))
    listed_at = None
    if other_info:
        match = _LISTED_YEAR.search(_fold(other_info))
        if match:
            listed_at = match.group(1)

    raw = {
        "source": "EOCN",
        "sheet": sheet_name,
        "entry_number": index,
        "classification": classification,
        "nationality": _text(row, _col(columns, H_NATIONALITY)),
        "birth_place": _text(row, _col(columns, H_BIRTH_PLACE)),
        "address": _prune({
            "street": _text(row, _col(columns, H_STREET)),
            "city": _text(row, _col(columns, H_CITY)),
            "country": _text(row, _col(columns, H_ADDR_COUNTRY)),
        }),
        "document_type": _text(row, _col(columns, H_DOC_TYPE)),
        "document_number": _text(row, _col(columns, H_DOC_NUMBER)),
        "document_issuer": _text(row, _col(columns, H_DOC_ISSUER)),
        "licence_number": _text(row, _col(columns, H_LICENCE_NO)),
        "headquarters": _text(row, _col(columns, H_HQ)),
        "notes": _text(row, _col(columns, H_NOTES)),
        # The Cabinet Resolution that made the designation. This is the
        # citation an operator needs in front of them before freezing
        # anything, so it is carried through to the evidence pack rather
        # than discarded at ingestion.
        "designation_basis": other_info,
    }

    return SourceEntity(
        source_id=f"AE-LTL-{kind}-{index}",
        schema_type="Person" if is_person else "Organization",
        caption=caption,
        names=names,
        countries=countries,
        birth_date=_birth_date(_at(row, _col(columns, H_BIRTH_DATE))),
        gender=None,          # not published in this list
        topics=["sanction", "crime.terror"],
        programs=[PROGRAM],
        identifiers=_identifiers(row, columns),
        listed_at=listed_at,
        raw=_prune(raw),
    )


# --------------------------------------------------------------------------
# Page scraping helpers
# --------------------------------------------------------------------------


def _candidate_file_ids(html: str) -> list[str]:
    """FileIDs published under the "Local Terrorist List" heading, Excel first.

    The page publishes many files -- decrees, guidance notes, the UN
    consolidated list, the control list -- all through the same
    `DownloadFile?FileID=` endpoint, so a link is only identifiable by what it
    sits underneath. This walks the page once, tracks which list heading is
    currently in scope, and keeps the links found under ours.

    Excel links are returned first because they are what `parse` can read; the
    PDF of the same list follows as a fallback that will be rejected on its
    magic bytes, which produces a precise error instead of a silent miss.
    """
    excel: list[str] = []
    other: list[str] = []
    heading_at: int | None = None
    pending_excel = False

    for match in _MARKER.finditer(html):
        if match.group("ltl"):
            heading_at = match.end()
            pending_excel = False
        elif match.group("other"):
            heading_at = None
            pending_excel = False
        elif match.group("excel"):
            pending_excel = True
        elif match.group("pdf"):
            pending_excel = False
        else:
            file_id = match.group("fid")
            in_scope = heading_at is not None and match.start() - heading_at <= _WINDOW_CHARS
            if in_scope:
                bucket = excel if pending_excel else other
                if file_id not in excel and file_id not in other:
                    bucket.append(file_id)
            pending_excel = False

    return excel + other


def _looks_like_workbook(payload: bytes) -> bool:
    return payload.startswith(_OLE2_MAGIC) or payload.startswith(_ZIP_MAGIC)


def _prune(values: dict) -> dict:
    """Drop empty keys so the stored raw record shows what the source had, not
    a wall of nulls in the evidence pack."""
    return {k: v for k, v in values.items() if v not in (None, "", {}, [])}


def _describe(payload: bytes) -> str:
    if payload.startswith(b"%PDF"):
        return "pdf"
    if payload.startswith(_OLE2_MAGIC):
        return "xls"
    if payload.startswith(_ZIP_MAGIC):
        return "xlsx"
    return "unknown"


def uae_local_terrorists() -> EOCNLocalTerroristListAdapter:
    """The mandatory UAE list, from the primary source."""
    return EOCNLocalTerroristListAdapter()
